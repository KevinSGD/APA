import io
import json
import re
from datetime import date
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


st.set_page_config(
    page_title="Agente Académico APA 7",
    page_icon="📄",
    layout="centered",
)

DEFAULT_STUDENT = "Kevin Steve Garzon Diaz"
DEFAULT_INSTITUTION = "Universidad Manuela Beltrán\nIngeniería de Software\nFacultad de Ingeniería\nColombia"
MODEL_NAME = "gemini-1.5-pro"
PORTADA_IMAGE = Path(__file__).with_name("Imagen1.jpg")


def clean_json_response(response_text: str) -> dict[str, Any]:
    """Extract and parse a JSON object, including responses wrapped in Markdown fences."""
    cleaned = re.sub(r"```(?:json)?", "", response_text, flags=re.IGNORECASE).strip()
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("Gemini no devolvió un objeto JSON válido.")
    return json.loads(cleaned[start : end + 1])


def get_model() -> Any | None:
    try:
        import google.generativeai as genai
    except ModuleNotFoundError:
        return None
    try:
        api_key = st.secrets.get("GEMINI_API_KEY", "")
    except (FileNotFoundError, KeyError):
        api_key = ""
    if not api_key:
        return None
    genai.configure(api_key=api_key)
    return genai.GenerativeModel(MODEL_NAME)


def fallback_plan(materia: str, titulo: str) -> list[dict[str, Any]]:
    return [
        {"nivel": 1, "subtitulo": "Introducción"},
        {"nivel": 1, "subtitulo": f"Fundamentos de {titulo}"},
        {"nivel": 2, "subtitulo": f"Conceptos centrales en {materia}"},
        {"nivel": 2, "subtitulo": "Aplicaciones y desafíos actuales"},
        {"nivel": 1, "subtitulo": "Conclusiones"},
    ]


def planificar_estructura(model: Any | None, materia: str, titulo: str) -> list[dict[str, Any]]:
    if model is None:
        return fallback_plan(materia, titulo)

    prompt = f"""
Actúa como agente planificador académico experto en APA 7. Descompón este trabajo en una
tabla de contenido lógica y progresiva:
Materia: {materia}
Título: {titulo}

Devuelve exclusivamente JSON válido con esta forma:
{{"estructura": [{{"nivel": 1, "subtitulo": "..."}}]}}

Incluye introducción, desarrollo suficiente con niveles APA 1, 2 y 3 cuando sea útil,
y conclusiones. No uses numeración dentro de los subtítulos. Entrega entre 5 y 12 elementos.
"""
    data = clean_json_response(model.generate_content(prompt).text)
    estructura = data.get("estructura", [])
    if not isinstance(estructura, list) or not estructura:
        raise ValueError("La planificación recibida no contiene una estructura válida.")
    return normalizar_estructura(estructura)


def normalizar_estructura(items: list[Any]) -> list[dict[str, Any]]:
    result = []
    for item in items:
        if isinstance(item, dict):
            nivel = item.get("nivel", 1)
            subtitulo = str(item.get("subtitulo", "")).strip()
        else:
            match = re.match(r"Nivel\s+(\d)\s*\|\s*(.+)", str(item), re.IGNORECASE)
            nivel = match.group(1) if match else 1
            subtitulo = match.group(2).strip() if match else str(item).strip()
        if subtitulo and str(nivel).isdigit() and 1 <= int(nivel) <= 3:
            result.append({"nivel": int(nivel), "subtitulo": subtitulo})
    if not result:
        raise ValueError("La estructura debe incluir subtítulos entre los niveles 1 y 3.")
    return result


def estructura_a_texto(estructura: list[dict[str, Any]]) -> str:
    return "\n".join(f"Nivel {item['nivel']} | {item['subtitulo']}" for item in estructura)


def texto_a_estructura(texto: str) -> list[dict[str, Any]]:
    return normalizar_estructura([line for line in texto.splitlines() if line.strip()])


def fallback_contenido(materia: str, titulo: str) -> tuple[list[str], list[dict[str, str]]]:
    parrafos = [
        f"El análisis de {titulo.lower()} en el contexto de {materia} requiere relacionar sus fundamentos conceptuales con los problemas que pretende resolver. Una perspectiva sistemática permite delimitar el fenómeno, identificar sus variables principales y valorar sus implicaciones académicas.",
        f"Desde una perspectiva aplicada, las decisiones relacionadas con {titulo.lower()} deben sustentarse en criterios verificables de calidad, contexto y sostenibilidad. La literatura especializada destaca que la claridad del problema y la trazabilidad de las decisiones fortalecen la validez de los resultados (Martin, 2017; Pressman & Maxim, 2020).",
        "La revisión de sus aplicaciones muestra que los resultados dependen tanto de los principios técnicos como de la capacidad de evaluar riesgos, restricciones y efectos a largo plazo. Por ello, el análisis debe combinar fundamentos teóricos, evidencia documentada y una lectura crítica de cada alternativa.",
        "En conclusión, comprender este tema permite formular soluciones más coherentes y defendibles. La integración de conceptos, evidencia y evaluación continua ofrece una base sólida para futuras investigaciones y para la práctica profesional.",
    ]
    referencias = [
        {"texto": "Martin, R. C. (2017). Clean architecture: A craftsman's guide to software structure and design. Pearson.", "tipo": "libro"},
        {"texto": "Pressman, R. S., & Maxim, B. R. (2020). Software engineering: A practitioner's approach (9th ed.). McGraw-Hill.", "tipo": "libro"},
    ]
    return parrafos, referencias


def redactar_contenido(
    model: Any | None,
    materia: str,
    titulo: str,
    estructura: list[dict[str, Any]],
) -> tuple[list[str], list[dict[str, Any]]]:
    if model is None:
        return fallback_contenido(materia, titulo)

    outline = estructura_a_texto(estructura)
    prompt = f"""
Actúa como agente redactor académico experto en APA 7. Redacta un trabajo formal, profundo,
coherente y completamente en español.
Materia: {materia}
Título: {titulo}
Estructura aprobada:
{outline}

Devuelve exclusivamente JSON válido con esta forma:
{{
  "parrafos": ["..."],
  "referencias": [
    {{"texto": "Referencia completa en APA 7", "tipo": "libro|articulo|repositorio"}}
  ]
}}

Reglas: desarrolla introducción, análisis y conclusiones; incorpora citas parentéticas o
narrativas APA 7 dentro de los párrafos; respalda cada afirmación técnica importante con
una fuente de la lista final; no inventes DOI, URL ni datos editoriales; usa solo referencias
reales o claramente coherentes con el tema; incluye al menos una referencia por afirmación
técnica importante. Las referencias deben corresponder al tipo indicado.
"""
    data = clean_json_response(model.generate_content(prompt).text)
    parrafos = data.get("parrafos", [])
    referencias = data.get("referencias", [])
    if not isinstance(parrafos, list) or not all(isinstance(item, str) for item in parrafos):
        raise ValueError("La respuesta no contiene una lista válida de párrafos.")
    if not parrafos or not isinstance(referencias, list) or not referencias:
        raise ValueError("La respuesta académica debe incluir párrafos y referencias.")
    return [item.strip() for item in parrafos if item.strip()], normalizar_referencias(referencias)


def normalizar_referencias(referencias: list[Any]) -> list[dict[str, str]]:
    result = []
    for referencia in referencias:
        if isinstance(referencia, dict):
            texto = str(referencia.get("texto", "")).strip()
            tipo = str(referencia.get("tipo", "libro")).lower().strip()
        else:
            texto = str(referencia).strip()
            tipo = clasificar_referencia(texto)
        if texto:
            result.append({"texto": texto, "tipo": tipo if tipo in {"libro", "articulo", "repositorio"} else "libro"})
    if not result:
        raise ValueError("No se encontraron referencias utilizables.")
    return result


def clasificar_referencia(texto: str) -> str:
    lowered = texto.lower()
    if any(host in lowered for host in ("github.com", "gitlab.com", "bitbucket.org")):
        return "repositorio"
    if "doi.org" in lowered or re.search(r"\d+\s*\(\d+\)", texto):
        return "articulo"
    return "libro"


def validar_referencia_apa(referencia: dict[str, str]) -> list[str]:
    texto = referencia["texto"]
    tipo = referencia["tipo"]
    errores = []
    if not re.search(r"\((?:19|20)\d{2}|n\.d\.\)", texto, re.IGNORECASE):
        errores.append("falta el año entre paréntesis")
    if "." not in texto:
        errores.append("faltan separadores de autor y título")
    if tipo == "articulo" and not ("doi.org" in texto.lower() or re.search(r"\d+\s*\(\d+\)", texto)):
        errores.append("un artículo debe incluir volumen/número o DOI")
    if tipo == "repositorio" and not re.search(r"https?://", texto, re.IGNORECASE):
        errores.append("un repositorio debe incluir una URL")
    return errores


def validar_referencias(referencias: list[dict[str, str]]) -> tuple[list[dict[str, str]], list[str]]:
    ordenadas = sorted(referencias, key=lambda item: item["texto"].casefold())
    advertencias = []
    for referencia in ordenadas:
        errores = validar_referencia_apa(referencia)
        if errores:
            advertencias.append(f"{referencia['texto']} ({'; '.join(errores)}).")
    return ordenadas, advertencias


def revisar_citas(parrafos: list[str]) -> list[str]:
    citation_pattern = re.compile(r"\([A-ZÁÉÍÓÚÑ][^()]{1,80},\s*(?:19|20)\d{2}\)")
    return [
        f"El párrafo {index} no contiene una cita APA 7 identificable."
        for index, parrafo in enumerate(parrafos, start=1)
        if not citation_pattern.search(parrafo)
    ]


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        add_page_number(section)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 2.0
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_page_number(section: Any) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def add_centered_text(doc: Document, text: str, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_cover_image(doc: Document) -> None:
    if not PORTADA_IMAGE.is_file():
        raise FileNotFoundError(
            f"No se encontró la imagen de portada en: {PORTADA_IMAGE.name}"
        )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run()
    run.add_picture(str(PORTADA_IMAGE), width=Inches(1.25))


def add_cover_lines(
    doc: Document,
    text: str,
    bold: bool = False,
    space_before: Pt | None = None,
) -> None:
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.space_before = space_before if index == 0 and space_before else Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_apa_heading(doc: Document, text: str, level: int) -> None:
    """Add a real Word heading using the APA 7 presentation for levels 1-3."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.left_indent = Inches(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = True
    run.italic = level == 3
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    paragraph.style = doc.styles[f"Heading {min(max(level, 1), 3)}"]
    paragraph.style.font.name = "Times New Roman"
    paragraph.style.font.size = Pt(12)
    paragraph.style.font.bold = True
    paragraph.style.paragraph_format.line_spacing = 2.0
    paragraph.style.paragraph_format.space_before = Pt(0)
    paragraph.style.paragraph_format.space_after = Pt(0)


def add_table_of_contents(doc: Document) -> None:
    add_centered_text(doc, "Tabla de contenido", bold=True)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Abra el documento en Word para actualizar la tabla de contenido."
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_separate, placeholder, field_end))


def ensure_introduction_first(estructura: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized = normalizar_estructura(estructura)
    without_introduction = [
        item for item in normalized if item["subtitulo"].casefold() != "introducción"
    ]
    return [{"nivel": 1, "subtitulo": "Introducción"}, *without_introduction]


def spanish_date() -> str:
    months = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre")
    current = date.today()
    return f"{current.day} de {months[current.month - 1]} de {current.year}"


def build_document(
    titulo: str,
    materia: str,
    estudiante: str,
    institucion: str,
    docente: str,
    parrafos: list[str],
    referencias: list[dict[str, str]],
    estructura: list[dict[str, Any]] | None = None,
) -> io.BytesIO:
    doc = Document()
    configure_document(doc)

    for _ in range(1):
        doc.add_paragraph()
    add_cover_image(doc)
    add_cover_lines(doc, materia)
    add_cover_lines(doc, titulo, bold=True)
    add_cover_lines(doc, estudiante, space_before=Pt(96))
    add_cover_lines(doc, institucion, space_before=Pt(96))
    add_cover_lines(doc, str(date.today().year))
    doc.add_page_break()

    estructura = ensure_introduction_first(estructura or [
        {"nivel": 1, "subtitulo": "Conclusiones"},
    ])
    add_table_of_contents(doc)
    doc.add_page_break()

    add_centered_text(doc, titulo, bold=True)
    add_apa_heading(doc, "Introducción", 1)
    remaining_paragraphs = list(parrafos)
    for index, item in enumerate(estructura):
        if index == 0:
            heading_already_added = True
        else:
            add_apa_heading(doc, item["subtitulo"], item["nivel"])
            heading_already_added = False
        if remaining_paragraphs and (index == 0 or not heading_already_added):
            texto = remaining_paragraphs.pop(0)
            paragraph = doc.add_paragraph(texto)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Inches(0.5)
    for texto in remaining_paragraphs:
        paragraph = doc.add_paragraph(texto)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Inches(0.5)

    doc.add_page_break()
    add_apa_heading(doc, "Referencias", 1)
    for referencia in referencias:
        paragraph = doc.add_paragraph(referencia["texto"])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def safe_filename(title: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", title).strip(" .")
    return f"{(cleaned or 'actividad')[:100].replace(' ', '_')}_APA7.docx"


def iter_document_paragraphs(doc: Document) -> list[Any]:
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def document_text_snapshot(doc: Document) -> list[str]:
    """Capture textual content only; formatting and added APA fields are excluded."""
    snapshot = []
    for paragraph in iter_document_paragraphs(doc):
        text = paragraph.text
        if text.strip():
            snapshot.append(text)
    return snapshot


def original_paragraph_snapshot(doc: Document) -> list[tuple[Any, str]]:
    return [(paragraph, paragraph.text) for paragraph in iter_document_paragraphs(doc)]


def original_content_unchanged(snapshot: list[tuple[Any, str]]) -> bool:
    for paragraph, original_text in snapshot:
        if paragraph.text != original_text:
            return False
    return True


def is_subsequence(expected: list[str], actual: list[str]) -> bool:
    position = 0
    for item in actual:
        if position < len(expected) and item == expected[position]:
            position += 1
    return position == len(expected)


def all_document_runs(doc: Document) -> list[Any]:
    runs = []
    for paragraph in iter_document_paragraphs(doc):
        runs.extend(paragraph.runs)
    return runs


def has_page_field(section: Any) -> bool:
    return "PAGE" in section.header._element.xml.upper()


def analyze_uploaded_document(doc: Document) -> dict[str, Any]:
    paragraphs = [paragraph for paragraph in iter_document_paragraphs(doc) if paragraph.text.strip()]
    runs = all_document_runs(doc)
    normal_style = doc.styles["Normal"]
    style_font_ok = normal_style.font.name == "Times New Roman" and normal_style.font.size == Pt(12)
    section_margins_ok = all(
        abs(section.top_margin.inches - 1) < 0.01
        and abs(section.bottom_margin.inches - 1) < 0.01
        and abs(section.left_margin.inches - 1) < 0.01
        and abs(section.right_margin.inches - 1) < 0.01
        for section in doc.sections
    )
    run_font_ok = bool(runs) and all(
        run.font.name in (None, "Times New Roman") and run.font.size in (None, Pt(12))
        for run in runs
    )
    references_index = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip().casefold() == "referencias"),
        None,
    )
    reference_paragraphs = paragraphs[references_index + 1 :] if references_index is not None else []
    references_ok = bool(reference_paragraphs) and all(
        paragraph.paragraph_format.left_indent is not None
        and paragraph.paragraph_format.first_line_indent is not None
        and paragraph.paragraph_format.left_indent.inches >= 0.49
        and paragraph.paragraph_format.first_line_indent.inches <= -0.49
        for paragraph in reference_paragraphs
    )
    headings = [
        paragraph for paragraph in paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]
    citations = sum(
        len(re.findall(r"\([A-ZÁÉÍÓÚÑ][^()]{1,80},\s*(?:19|20)\d{2}\)", paragraph.text))
        for paragraph in paragraphs
    )
    captions = [
        paragraph for paragraph in paragraphs
        if re.search(r"^(Figura|Figure)\s+\d+", paragraph.text.strip(), re.IGNORECASE)
    ]
    checks = [
        ("Márgenes de 2,54 cm", section_margins_ok, "Todos los márgenes deben medir 1 pulgada."),
        ("Times New Roman de 12 puntos", style_font_ok and run_font_ok, "La fuente base o algunos textos no tienen Times New Roman de 12 puntos."),
        ("Interlineado y espaciado", normal_style.paragraph_format.line_spacing == 2.0 and normal_style.paragraph_format.space_after == Pt(0), "Se requiere interlineado doble y espaciado posterior de 0 puntos."),
        ("Sangría y alineación", bool(paragraphs), "Se revisará la sangría de los párrafos de contenido."),
        ("Encabezados APA", bool(headings), "No se detectaron estilos Heading 1, Heading 2 o Heading 3."),
        ("Referencias", references_ok, "No se detectó sangría francesa consistente en las referencias."),
        ("Citas dentro del texto", citations > 0, "No se detectaron citas parentéticas con año."),
        ("Numeración de páginas", all(has_page_field(section) for section in doc.sections), "No se detectó el campo PAGE en los encabezados."),
        ("Tabla de figuras", bool(captions), "No se detectaron pies de figura con el formato Figura N."),
    ]
    passed = sum(item[1] for item in checks)
    return {
        "score": round(passed / len(checks) * 100),
        "checks": [{"name": name, "ok": ok, "detail": detail} for name, ok, detail in checks],
        "headings": len(headings),
        "references": len(reference_paragraphs),
        "citations": citations,
        "captions": len(captions),
        "paragraphs": len(paragraphs),
    }


def set_run_font(run: Any) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def apply_selected_formatting(doc: Document, options: set[str]) -> None:
    if "margins" in options:
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
    if "typography" in options:
        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(12)
        for run in all_document_runs(doc):
            set_run_font(run)
    if "spacing" in options:
        normal_style = doc.styles["Normal"]
        normal_style.paragraph_format.line_spacing = 2.0
        normal_style.paragraph_format.space_before = Pt(0)
        normal_style.paragraph_format.space_after = Pt(0)
        for paragraph in iter_document_paragraphs(doc):
            paragraph.paragraph_format.line_spacing = 2.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
    if "alignment" in options or "indentation" in options:
        references_started = False
        for paragraph in iter_document_paragraphs(doc):
            if paragraph.text.strip().casefold() == "referencias":
                references_started = True
            if "alignment" in options:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if "indentation" in options and paragraph.text.strip() and not references_started:
                paragraph.paragraph_format.first_line_indent = Inches(0.5)
    if "headings" in options:
        for paragraph in iter_document_paragraphs(doc):
            if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
                level = int(paragraph.style.name[-1])
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Inches(0)
                for run in paragraph.runs:
                    set_run_font(run)
                    run.bold = True
                    run.italic = level == 3
    if "references" in options:
        references_started = False
        for paragraph in iter_document_paragraphs(doc):
            if paragraph.text.strip().casefold() == "referencias":
                references_started = True
                continue
            if references_started and paragraph.text.strip():
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.5)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 2.0
    if "page_numbers" in options:
        for section in doc.sections:
            if not has_page_field(section):
                add_page_number(section)


def prepend_generated_content(doc: Document, builder: Any) -> None:
    body = doc._element.body
    section_properties = next((child for child in body if child.tag == qn("w:sectPr")), None)
    original = [child for child in body if child is not section_properties]
    anchor = original[0] if original else section_properties
    before = set(original)
    builder()
    generated = [child for child in body if child is not section_properties and child not in before]
    for child in reversed(generated):
        body.insert(body.index(anchor), child)


def apply_cover_to_uploaded(doc: Document, estudiante: str, institucion: str, materia: str, titulo: str) -> None:
    def builder() -> None:
        doc.add_paragraph()
        add_cover_image(doc)
        add_cover_lines(doc, materia)
        add_cover_lines(doc, titulo, bold=True)
        add_cover_lines(doc, estudiante, space_before=Pt(96))
        add_cover_lines(doc, institucion, space_before=Pt(96))
        add_cover_lines(doc, str(date.today().year))
        doc.add_page_break()
    prepend_generated_content(doc, builder)


def apply_toc_to_uploaded(doc: Document) -> None:
    def builder() -> None:
        add_table_of_contents(doc)
        doc.add_page_break()
    prepend_generated_content(doc, builder)


def apply_figures_toc(doc: Document) -> None:
    def builder() -> None:
        add_centered_text(doc, "Tabla de figuras", bold=True)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        field_begin = OxmlElement("w:fldChar")
        field_begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = ' TOC \\c "Figura" \\h \\z \\u '
        field_separate = OxmlElement("w:fldChar")
        field_separate.set(qn("w:fldCharType"), "separate")
        placeholder = OxmlElement("w:t")
        placeholder.text = "Abra el documento en Word para actualizar la tabla de figuras."
        field_end = OxmlElement("w:fldChar")
        field_end.set(qn("w:fldCharType"), "end")
        run._r.extend((field_begin, instruction, field_separate, placeholder, field_end))
        doc.add_page_break()
    prepend_generated_content(doc, builder)


def analyze_and_format_uploaded(uploaded_file: Any, options: set[str], cover_data: dict[str, str]) -> tuple[io.BytesIO, dict[str, Any]]:
    doc = Document(uploaded_file)
    before = document_text_snapshot(doc)
    original_elements = original_paragraph_snapshot(doc)
    if options & {"margins", "typography", "spacing", "alignment", "indentation", "headings", "references", "page_numbers"}:
        apply_selected_formatting(doc, options)
    if "cover" in options:
        apply_cover_to_uploaded(doc, cover_data["estudiante"], cover_data["institucion"], cover_data["materia"], cover_data["titulo"])
    if "toc" in options:
        apply_toc_to_uploaded(doc)
    if "figures_toc" in options:
        apply_figures_toc(doc)
    after = document_text_snapshot(doc)
    if not is_subsequence(before, after) or not original_content_unchanged(original_elements):
        raise ValueError("La verificación detectó una alteración del contenido textual original; no se exportó el archivo.")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer, analyze_uploaded_document(doc)


def update_progress(progress: Any, status: Any, value: int, message: str) -> None:
    progress.progress(value)
    status.update(label=message, state="running")


st.title("📚 Agente Autónomo de Investigación y Generación APA 7")
st.markdown("Crea documentos nuevos o revisa un Word existente sin alterar su contenido textual.")
mode = st.radio(
    "Modo de trabajo",
    ["Crear documento nuevo", "Revisar documento Word"],
    horizontal=True,
)

if mode == "Crear documento nuevo":
    with st.form("apa_form"):
        materia = st.text_input("Materia / Asignatura", placeholder="Ej. Arquitectura de Software")
        titulo_actividad = st.text_input("Título de la Actividad", placeholder="Ej. Diseño de Sistemas Distribuidos")
        nombre_estudiante = st.text_area("Nombres de los integrantes", value=DEFAULT_STUDENT, height=90)
        institucion = st.text_area("Universidad, programa, facultad y país", value=DEFAULT_INSTITUTION, height=120)
        nombre_profesor = st.text_input("Nombre del Profesor / Docente", placeholder="Ej. Ing. Carlos Pérez")
        plan_submitted = st.form_submit_button("1. Planificar estructura")

    if plan_submitted:
        if not materia.strip() or not titulo_actividad.strip():
            st.warning("Completa la materia y el título de la actividad para planificar el documento.")
        else:
            try:
                with st.spinner("🧭 El agente planificador está estructurando el documento..."):
                    model = get_model()
                    estructura = planificar_estructura(model, materia.strip(), titulo_actividad.strip())
                    st.session_state["apa_context"] = {
                        "materia": materia.strip(),
                        "titulo": titulo_actividad.strip(),
                        "estudiante": nombre_estudiante.strip() or DEFAULT_STUDENT,
                        "institucion": institucion.strip() or DEFAULT_INSTITUTION,
                        "docente": nombre_profesor.strip(),
                    }
                    st.session_state["apa_plan"] = estructura_a_texto(estructura)
                st.success("Estructura lista para revisión.")
            except Exception as error:
                st.error(f"No fue posible planificar el documento: {error}")

    context = st.session_state.get("apa_context")
    if context and "apa_plan" in st.session_state:
        with st.expander("Revisar y ajustar estructura APA", expanded=True):
            st.caption("Una línea por subtítulo, con el formato: Nivel 1 | Subtítulo. Se admiten niveles 1, 2 y 3.")
            st.session_state["apa_plan"] = st.text_area("Tabla de contenido", value=st.session_state["apa_plan"], height=180)
            generate_submitted = st.button("2. Redactar, validar y generar documento", type="primary")

        if generate_submitted:
            try:
                estructura = texto_a_estructura(st.session_state["apa_plan"])
                progress = st.progress(0)
                status = st.status("Inicializando agente autónomo...", expanded=True)
                model = get_model()

                update_progress(progress, status, 25, "Planificando estructura del documento...")
                update_progress(progress, status, 50, "Redactando contenido académico e incorporando citas...")
                parrafos, referencias = redactar_contenido(model, context["materia"], context["titulo"], estructura)

                update_progress(progress, status, 75, "Validando referencias bibliográficas y formato APA...")
                referencias, reference_warnings = validar_referencias(referencias)
                citation_warnings = revisar_citas(parrafos)
                if reference_warnings or citation_warnings:
                    with st.expander("Revisión académica", expanded=False):
                        for warning in reference_warnings + citation_warnings:
                            st.warning(warning)

                update_progress(progress, status, 90, "Compilando documento final en Word...")
                document_buffer = build_document(
                    context["titulo"],
                    context["materia"],
                    context["estudiante"],
                    context["institucion"],
                    context["docente"],
                    parrafos,
                    referencias,
                    estructura,
                )
                progress.progress(100)
                status.update(label="Documento listo para descargar.", state="complete")
                st.success("Documento generado con flujo Plan-Execute-Review.")
                st.download_button(
                    label="📥 Descargar Word (.docx)",
                    data=document_buffer,
                    file_name=safe_filename(context["titulo"]),
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            except Exception as error:
                st.error(f"Ocurrió un error durante la generación: {error}")
else:
    st.subheader("Revisar y aplicar APA 7 a un documento existente")
    st.info("El texto original se conserva. Solo se aplican los componentes que selecciones.")
    uploaded_file = st.file_uploader("Sube un documento Word", type=["docx"])
    if uploaded_file is not None:
        try:
            uploaded_file.seek(0)
            uploaded_doc = Document(uploaded_file)
            report = analyze_uploaded_document(uploaded_doc)
            st.metric("Puntuación APA estimada", f"{report['score']} / 100")
            st.caption(
                f"Párrafos: {report['paragraphs']} · Encabezados: {report['headings']} · "
                f"Referencias: {report['references']} · Citas detectadas: {report['citations']} · "
                f"Figuras: {report['captions']}"
            )
            with st.expander("Ver análisis y recomendaciones", expanded=True):
                for check in report["checks"]:
                    icon = "✅" if check["ok"] else "⚠️"
                    st.markdown(f"{icon} **{check['name']}**: {check['detail'] if not check['ok'] else 'Correcto.'}")

            option_labels = {
                "margins": "Márgenes APA (2,54 cm)",
                "typography": "Times New Roman de 12 puntos",
                "spacing": "Interlineado doble y espaciado de párrafos",
                "alignment": "Alineación a la izquierda",
                "indentation": "Sangría de primera línea",
                "headings": "Estilos APA para encabezados existentes",
                "references": "Sangría francesa en referencias",
                "page_numbers": "Numeración de páginas",
                "cover": "Agregar portada APA 7",
                "toc": "Agregar tabla de contenido",
                "figures_toc": "Agregar tabla de figuras",
            }
            selected_labels = st.multiselect(
                "Selecciona las correcciones o componentes que deseas aplicar",
                list(option_labels.values()),
                default=[option_labels[key] for key in ("margins", "typography", "spacing", "alignment", "indentation", "references", "page_numbers")],
            )
            selected = {key for key, label in option_labels.items() if label in selected_labels}

            cover_data = {
                "materia": st.text_input("Materia para la portada", value=""),
                "titulo": st.text_input("Título para la portada", value=""),
                "estudiante": st.text_area("Integrantes para la portada", value=DEFAULT_STUDENT, height=80),
                "institucion": st.text_area("Datos institucionales para la portada", value=DEFAULT_INSTITUTION, height=110),
            }
            if st.button("Aplicar selección y descargar copia APA", type="primary"):
                uploaded_file.seek(0)
                with st.spinner("Analizando y aplicando únicamente los componentes seleccionados..."):
                    output, final_report = analyze_and_format_uploaded(uploaded_file, selected, cover_data)
                st.success("Documento revisado sin alterar el contenido textual original.")
                st.download_button(
                    "📥 Descargar copia revisada",
                    data=output,
                    file_name=safe_filename(Path(uploaded_file.name).stem + "_APA7_revisado"),
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                st.caption(f"Puntuación después del formato: {final_report['score']} / 100")
        except Exception as error:
            st.error(f"No fue posible analizar el documento: {error}")

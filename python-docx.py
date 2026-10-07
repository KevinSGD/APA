from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def crear_documento_apa(materia, titulo_actividad, contenido_generado, referencias):
    doc = Document()

    # 1. Configurar Márgenes (2.54 cm = 1 pulgada)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 2. Configurar Estilo Normal (Times New Roman 12pt, Interlineado Doble)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 2.0
    paragraph_format.space_after = Pt(0)
    paragraph_format.space_before = Pt(0)

    # --- PORTADA (Formato Estudiante APA 7) ---
    # Espacios en blanco superiores para centrar la portada verticalmente
    for _ in range(3):
        doc.add_paragraph()

    # Título del trabajo
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = p_titulo.add_run(titulo_actividad)
    run_titulo.bold = True
    run_titulo.font.size = Pt(12)

    doc.add_paragraph() # Espacio

    # Datos del estudiante e institución
    p_datos = doc.add_paragraph()
    p_datos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_datos.add_run("Nombre del Estudiante\n")
    p_datos.add_run("Facultad de Ingeniería, Universidad\n")
    p_datos.add_run(f"{materia}\n")
    p_datos.add_run("Nombre del Profesor\n")
    p_datos.add_run("Fecha de Entrega")

    # Salto de página obligatorio tras la portada
    doc.add_page_break()

    # --- CUERPO DEL TRABAJO ---
    # Título de la primera sección (repite el título de la actividad según APA 7)
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_h = p_head.add_run(titulo_actividad)
    run_h.bold = True

    # Agregar el contenido redactado por el LLM con sangría de primera línea (1.27 cm)
    for parrafo_texto in contenido_generado:
        p = doc.add_paragraph(parrafo_texto)
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # --- REFERENCIAS ---
    doc.add_page_break()
    p_ref_title = doc.add_paragraph()
    p_ref_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_rt = p_ref_title.add_run("Referencias")
    run_rt.bold = True

    for ref in referencias:
        p_ref = doc.add_paragraph(ref)
        # Sangría francesa para referencias APA
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)
        p_ref.paragraph_format.line_spacing = 2.0

    # Guardar el archivo
    nombre_archivo = f"{titulo_actividad.replace(' ', '_')}_APA7.docx"
    doc.save(nombre_archivo)
    return nombre_archivo

# Ejemplo de uso simulado
contenido_ejemplo = [
    "La ingeniería de software moderna requiere la implementación de arquitecturas limpias y patrones de diseño robustos que garanticen la mantenibilidad y escalabilidad de los sistemas a largo plazo.",
    "Durante el desarrollo de aplicaciones empresariales, el desacoplamiento de componentes permite mitigar riesgos técnicos y facilita las pruebas unitarias y de integración en entornos distribuidos."
]
referencias_ejemplo = [
    "Gamma, E., Helm, R., Johnson, R., & Vlissides, J. (1994). Design patterns: elements of reusable object-oriented software. Addison-Wesley.",
    "Martin, R. C. (2017). Clean architecture: a craftsman's guide to software structure and design. Prentice Hall."
]

# Ejecución de prueba
# crear_documento_apa("Arquitectura de Software", "Diseño de Sistemas Distribuidos", contenido_ejemplo, referencias_ejemplo)